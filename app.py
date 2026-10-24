import io
import os
import tempfile
from docx import Document
from PIL import Image
import pytesseract
import streamlit as st
import whisper

st.set_page_config(page_title="Media to Text Converter", page_icon="📝")
st.title("📝 Media to Text Converter")
st.write("Convert **Audio** (`.mp3`, `.opus`, `.wav`, `.m4a`, `.ogg`, `.flac`) or **Images** (`.jpg`, `.jpeg`, `.png`) to text.")

AUDIO_EXTENSIONS = ["opus", "mp3", "wav", "m4a", "ogg", "flac"]
IMAGE_EXTENSIONS = ["jpg", "jpeg", "png"]

@st.cache_resource
def load_whisper_model():
    return whisper.load_model("base")

def create_docx(text: str, title: str) -> io.BytesIO:
    """Creates a downloadable .docx file in memory."""
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(text)
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

uploaded_file = st.file_uploader(
    "Upload a file", 
    type=AUDIO_EXTENSIONS + IMAGE_EXTENSIONS
)

if uploaded_file is not None:
    file_ext = uploaded_file.name.split(".")[-1].lower()
    base_name = uploaded_file.name.rsplit(".", 1)[0]

    # --- CASE 1: IMAGE FILE (JPG, JPEG, PNG) ---
    if file_ext in IMAGE_EXTENSIONS:
        image = Image.open(uploaded_file)
        st.image(image, caption="Uploaded Image", use_container_width=True)

        if st.button("Extract Text from Image", type="primary"):
            with st.spinner("Extracting text with OCR..."):
                try:
                    extracted_text = pytesseract.image_to_string(image).strip()

                    if extracted_text:
                        st.success("Extraction complete!")
                        st.text_area("Extracted Text", value=extracted_text, height=250)

                        col1, col2 = st.columns(2)
                        with col1:
                            docx_file = create_docx(extracted_text, f"Transcription for {uploaded_file.name}")
                            st.download_button(
                                label="📥 Download Word (.docx)",
                                data=docx_file,
                                file_name=f"{base_name}.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            )
                        with col2:
                            st.download_button(
                                label="📄 Download Text (.txt)",
                                data=extracted_text,
                                file_name=f"{base_name}.txt",
                                mime="text/plain",
                            )
                    else:
                        st.warning("No readable text found in this image.")
                except Exception as e:
                    st.error(f"Error extracting text: {e}")

    # --- CASE 2: AUDIO FILE (MP3, OPUS, WAV, etc.) ---
    elif file_ext in AUDIO_EXTENSIONS:
        st.audio(uploaded_file)

        if st.button("Transcribe Audio", type="primary"):
            with st.spinner("Transcribing audio..."):
                model = load_whisper_model()

                with tempfile.NamedTemporaryFile(delete=False, suffix=f".{file_ext}") as temp_audio:
                    temp_audio.write(uploaded_file.read())
                    temp_audio_path = temp_audio.name

                try:
                    result = model.transcribe(temp_audio_path)
                    transcription = result["text"].strip()
                    lang = result.get("language", "unknown")

                    st.success("Transcription complete!")
                    st.caption(f"Detected language: **{lang.upper()}**")
                    st.text_area("Extracted Text", value=transcription, height=250)

                    col1, col2 = st.columns(2)
                    with col1:
                        docx_file = create_docx(transcription, f"Transcription for {uploaded_file.name}")
                        st.download_button(
                            label="📥 Download Word (.docx)",
                            data=docx_file,
                            file_name=f"{base_name}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    with col2:
                        st.download_button(
                            label="📄 Download Text (.txt)",
                            data=transcription,
                            file_name=f"{base_name}.txt",
                            mime="text/plain",
                        )
                except Exception as e:
                    st.error(f"Transcription failed: {e}")
                finally:
                    if os.path.exists(temp_audio_path):
                        os.remove(temp_audio_path)
